"""Unit and integration tests for FastAPI web application."""

import io
import json
import docx
import pytest
from fastapi.testclient import TestClient

from app import app


@pytest.fixture(scope="module")
def client():
    """Create a test client for the FastAPI application."""
    return TestClient(app)


def test_health_endpoint(client):
    """Health check endpoint returns status ok without document processing overhead."""
    response = client.get("/health")
    assert response.status_code == 200
    assert response.json() == {"status": "ok"}


def test_index_page(client):
    """Index page serves HTML with the upload form."""
    response = client.get("/")
    assert response.status_code == 200
    assert "text/html" in response.headers.get("content-type", "")
    assert "PII Redaction" in response.text
    assert "Redact Document" in response.text


def test_reject_non_docx_file(client):
    """Non-docx files are rejected with 400."""
    response = client.post(
        "/redact",
        files={"file": ("sample.pdf", b"%PDF-1.4 dummy", "application/pdf")},
    )
    assert response.status_code == 400
    assert "Only Microsoft Word (.docx) documents are supported" in response.json().get("detail", "")


def test_reject_empty_file(client):
    """Empty files are rejected with 400."""
    response = client.post(
        "/redact",
        files={"file": ("empty.docx", b"", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    assert "empty" in response.json().get("detail", "").lower()


def test_reject_corrupt_file(client):
    """Files with .docx extension but invalid zip headers are rejected with 400."""
    response = client.post(
        "/redact",
        files={"file": ("corrupt.docx", b"not a zip file", "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert response.status_code == 400
    assert "Invalid .docx format" in response.json().get("detail", "")


def test_redact_valid_docx(client):
    """Valid DOCX document is processed, redacted, and returned as a downloadable document."""
    doc = docx.Document()
    doc.add_paragraph("Director Rohit Kushal Hegde can be reached at cs.connect@kshinternational.com.")
    buf = io.BytesIO()
    doc.save(buf)
    docx_bytes = buf.getvalue()

    response = client.post(
        "/redact",
        files={"file": ("test_doc.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert response.status_code == 200
    assert "attachment" in response.headers.get("content-disposition", "")
    assert "Redacted_test_doc.docx" in response.headers.get("content-disposition", "")

    total_redactions = int(response.headers.get("X-Total-Redactions", "0"))
    assert total_redactions >= 2

    counts = json.loads(response.headers.get("X-Category-Counts", "{}"))
    assert "PERSON" in counts
    assert "EMAIL" in counts

    # Verify that the returned file is valid DOCX and sensitive values were redacted
    output_doc = docx.Document(io.BytesIO(response.content))
    output_text = "\n".join(p.text for p in output_doc.paragraphs)
    assert "Rohit Kushal Hegde" not in output_text
    assert "cs.connect@kshinternational.com" not in output_text
